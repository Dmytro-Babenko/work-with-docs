from setting_class import Setting, DATA_TEMPLATE_SETTINGS, GRAPH_SYMBOL
from docxtpl import DocxTemplate, InlineImage
from docx.shared import Cm
from pathlib import Path


class TemplateSetting(Setting):
    def __init__(self, setting_name, data) -> None:
        super().__init__(setting_name, data)

    def exe_program(self):
        doc_path = self.data['template word'].get_value()
        temp_image_folder = self.data['folder with tamplates'].get_value()
        doc = DocxTemplate(doc_path)

        placeholders = {}
        i = 1
        for image in temp_image_folder.iterdir():
            if image.suffix == '.png':
                placeholder = InlineImage(doc, str(image.absolute()), Cm(5))
                place = f'{GRAPH_SYMBOL}{i}'
                placeholders[place] = placeholder
                i += 1
        
        print(placeholders)
        doc.render(placeholders)
        doc.save(doc_path)
        pass
    
def main():
    main_settings = TemplateSetting('Main settings', DATA_TEMPLATE_SETTINGS)
    main_root = main_settings.make_setting_root()
    main_settings.make_fields(main_root)
    main_settings.key_element.tk_variable.trace_add('write', main_settings.update)
    main_root.mainloop()

if __name__ == '__main__':
    main()